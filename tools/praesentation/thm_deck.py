"""Gemeinsame Bausteine fuer die drei Praesentationen (THM-Vorlage FB M, Breitbild).

Gestaltung bewusst schlicht: Aussage als Folientitel, ein Bild oder Diagramm
je Folie, normaler Fliesstext. Keine Kennzahl-Kacheln.
"""
import copy
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

REPO = Path(__file__).resolve().parents[2]
QUELLEN = REPO / "praesentation" / "quellen"

GRUEN = RGBColor(0x80, 0xBA, 0x24)      # THM-Gruen
GRAU = RGBColor(0x4A, 0x5C, 0x66)       # THM-Grau (Fusszeile)
HELLGRAU = RGBColor(0xA9, 0xB4, 0xBA)
HELL = RGBColor(0xEE, 0xF1, 0xF2)
DUNKEL = RGBColor(0x26, 0x32, 0x38)
ROT = RGBColor(0xB3, 0x26, 0x3E)
WEISS = RGBColor(0xFF, 0xFF, 0xFF)
LINIE = RGBColor(0xDD, 0xDD, 0xDD)

A = "http://schemas.openxmlformats.org/drawingml/2006/main"


class Deck:
    def __init__(self, fusszeile="Aly Elbermawy"):
        src = QUELLEN / "THM_Vorlage.pptx"
        tmp = QUELLEN / "_vorlage_tmp.pptx"
        with zipfile.ZipFile(src) as zi, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zo:
            for it in zi.infolist():
                data = zi.read(it.filename)
                if it.filename.endswith(".xml") and b"Referenten" in data:
                    data = data.replace("Name des  Referenten".encode(), fusszeile.encode())
                zo.writestr(it, data)
        self.prs = Presentation(tmp)
        tmp.unlink()
        ids = self.prs.slides._sldIdLst
        for s in list(ids):
            self.prs.part.drop_rel(s.rId)
            ids.remove(s)
        self.L = {l.name: l for l in self.prs.slide_layouts}

    # ------------------------------------------------------------------ Folien
    @staticmethod
    def _ph(slide, idx):
        for p in slide.placeholders:
            if p.placeholder_format.idx == idx:
                return p
        return None

    def titelfolie(self, zeilen):
        """zeilen: Liste (text, groesse, fett, farbe)."""
        s = self.prs.slides.add_slide(self.L["1_THM-Titelfolie"])
        tf = self._ph(s, 10).text_frame
        tf.text = ""
        for i, (t, sz, b, c) in enumerate(zeilen):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(6 if i == 0 else 2)
            r = p.add_run()
            r.text = t
            r.font.size, r.font.bold, r.font.color.rgb = Pt(sz), b, c
        return s

    def folie(self, titel):
        s = self.prs.slides.add_slide(self.L["Textfolie"])
        self._ph(s, 0).text = titel
        return s

    # --------------------------------------------------------------- Elemente
    @staticmethod
    def text(slide, x, y, w, h, absaetze, size=18, color=DUNKEL, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, punkte=False, abstand=8):
        """absaetze: str oder Liste; ein Absatz ist str oder Liste von (text, optionen)."""
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0)
        tf.margin_top = tf.margin_bottom = Inches(0)
        tf.vertical_anchor = anchor
        liste = absaetze if isinstance(absaetze, list) else [absaetze]
        for i, ab in enumerate(liste):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.space_after = Pt(abstand)
            teile = ab if isinstance(ab, list) else [(ab, {})]
            if punkte:
                pPr = p._p.get_or_add_pPr()
                pPr.set("marL", "285750")
                pPr.set("indent", "-285750")
                bc = pPr.makeelement(f"{{{A}}}buClr", {})
                sc = bc.makeelement(f"{{{A}}}srgbClr", {"val": "80BA24"})
                bc.append(sc)
                pPr.append(bc)
                pPr.append(pPr.makeelement(f"{{{A}}}buChar", {"char": "–"}))
            for t, o in teile:
                r = p.add_run()
                r.text = t
                f = r.font
                f.size = Pt(o.get("size", size))
                f.bold = o.get("bold", bold)
                f.italic = o.get("italic", False)
                f.color.rgb = o.get("color", color)
        return tb

    @staticmethod
    def flaeche(slide, x, y, w, h, fill=HELL, line=None, form=MSO_SHAPE.RECTANGLE):
        sh = slide.shapes.add_shape(form, Inches(x), Inches(y), Inches(w), Inches(h))
        if fill is None:
            sh.fill.background()
        else:
            sh.fill.solid()
            sh.fill.fore_color.rgb = fill
        if line is None:
            sh.line.fill.background()
        else:
            sh.line.color.rgb = line
            sh.line.width = Pt(1.25)
        sh.shadow.inherit = False
        return sh

    @staticmethod
    def linie(slide, x1, y1, x2, y2, farbe=HELLGRAU, breite=1.5, pfeil=False):
        ln = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        ln.line.color.rgb = farbe
        ln.line.width = Pt(breite)
        if pfeil:
            le = ln.line._get_or_add_ln()
            le.append(le.makeelement(f"{{{A}}}tailEnd", {"type": "triangle", "w": "med", "len": "med"}))
        return ln

    @staticmethod
    def bild(slide, datei, x, y, w=None, h=None):
        kw = {}
        if w:
            kw["width"] = Inches(w)
        if h:
            kw["height"] = Inches(h)
        return slide.shapes.add_picture(str(datei), Inches(x), Inches(y), **kw)

    def bildunterschrift(self, slide, x, y, w, t):
        self.text(slide, x, y, w, 0.35, t, size=12, color=GRAU)

    @staticmethod
    def tabelle(slide, x, y, w, zeilen, breiten, size=15, h_zeile=0.42, hervor=(), zentriert_ab=1):
        rows, cols = len(zeilen), len(zeilen[0])
        t = slide.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w),
                                   Inches(h_zeile * rows)).table
        tblPr = t._tbl.tblPr
        for k in ("bandRow", "firstRow"):
            tblPr.set(k, "0")
        for j, b in enumerate(breiten):
            t.columns[j].width = Inches(b)
        for i, z in enumerate(zeilen):
            t.rows[i].height = Inches(h_zeile)
            for j, v in enumerate(z):
                c = t.cell(i, j)
                c.text = ""
                c.margin_left = c.margin_right = Inches(0.1)
                c.margin_top = c.margin_bottom = Inches(0.03)
                c.vertical_anchor = MSO_ANCHOR.MIDDLE
                p = c.text_frame.paragraphs[0]
                p.alignment = PP_ALIGN.LEFT if j < zentriert_ab else PP_ALIGN.CENTER
                r = p.add_run()
                r.text = str(v)
                r.font.size = Pt(size)
                r.font.bold = (i == 0) or (i in hervor)
                c.fill.solid()
                if i == 0:
                    c.fill.fore_color.rgb = GRAU
                    r.font.color.rgb = WEISS
                else:
                    c.fill.fore_color.rgb = RGBColor(0xEE, 0xF6, 0xE2) if i in hervor else (
                        WEISS if i % 2 else RGBColor(0xF4, 0xF6, 0xF7))
                    r.font.color.rgb = DUNKEL
        return t

    @staticmethod
    def diagramm(slide, x, y, w, h, kategorien, reihen, farben, titel=None, max_=None,
                 fmt='0"%"', legende=True, size=13, quer=False, beschriftung=True):
        cd = CategoryChartData()
        cd.categories = kategorien
        for name, werte in reihen:
            cd.add_series(name, werte)
        typ = XL_CHART_TYPE.BAR_CLUSTERED if quer else XL_CHART_TYPE.COLUMN_CLUSTERED
        ch = slide.shapes.add_chart(typ, Inches(x), Inches(y), Inches(w), Inches(h), cd).chart
        ch.font.size = Pt(size)
        ch.font.color.rgb = GRAU
        ch.has_title = bool(titel)
        if titel:
            ch.chart_title.text_frame.text = titel
            tr = ch.chart_title.text_frame.paragraphs[0].runs[0]
            tr.font.size, tr.font.bold, tr.font.color.rgb = Pt(size + 1), True, DUNKEL
        ch.has_legend = legende and len(reihen) > 1
        if ch.has_legend:
            ch.legend.position = XL_LEGEND_POSITION.TOP
            ch.legend.include_in_layout = False
            ch.legend.font.size = Pt(size)
        va = ch.value_axis
        va.has_major_gridlines = True
        va.major_gridlines.format.line.color.rgb = LINIE
        va.format.line.fill.background()
        va.tick_labels.font.size = Pt(size - 1)
        va.minimum_scale = 0
        if max_:
            va.maximum_scale = max_
        va.tick_labels.number_format = fmt
        va.tick_labels.number_format_is_linked = False
        ca = ch.category_axis
        ca.tick_labels.font.size = Pt(size)
        ca.format.line.color.rgb = RGBColor(0xBB, 0xBB, 0xBB)
        if quer:
            ca.reverse_order = True
            va.tick_labels.font.size = Pt(size - 2)
        pl = ch.plots[0]
        pl.gap_width = 70
        pl.overlap = -5
        if beschriftung:
            pl.has_data_labels = True
            dl = pl.data_labels
            dl.font.size = Pt(size - 1)
            dl.font.color.rgb = DUNKEL
            dl.number_format = fmt
            dl.number_format_is_linked = False
            dl.position = XL_LABEL_POSITION.OUTSIDE_END
        for s, f in zip(pl.series, farben):
            s.format.fill.solid()
            s.format.fill.fore_color.rgb = f
        return ch

    def ablauf(self, slide, x, y, w, schritte, h=0.95, hervor=(), size=14, unter=None):
        """Schlichte Kette aus Kaestchen mit Pfeilen; schritte: Liste (kopf, unterzeile)."""
        n = len(schritte)
        gap = 0.32
        bw = (w - gap * (n - 1)) / n
        for i, (kopf, zeile) in enumerate(schritte):
            xi = x + i * (bw + gap)
            hi = i in hervor
            self.flaeche(slide, xi, y, bw, h, fill=GRUEN if hi else HELL)
            self.text(slide, xi + 0.08, y + 0.12, bw - 0.16, 0.4, [[(kopf, {"bold": True})]],
                      size=size, color=WEISS if hi else DUNKEL, align=PP_ALIGN.CENTER)
            if zeile:
                self.text(slide, xi + 0.08, y + 0.5, bw - 0.16, h - 0.55, zeile.split("\n"),
                          size=size - 3, color=WEISS if hi else GRAU, align=PP_ALIGN.CENTER, abstand=0)
            if i < n - 1:
                self.linie(slide, xi + bw + 0.04, y + h / 2, xi + bw + gap - 0.04, y + h / 2,
                           farbe=GRAU, breite=1.75, pfeil=True)

    @staticmethod
    def notiz(slide, t):
        slide.notes_slide.notes_text_frame.text = t

    # ------------------------------------------------------------- Abschluss
    def speichern(self, datei):
        for sl in self.prs.slides:
            for p in list(sl.placeholders):
                if p.has_text_frame and not p.text_frame.text.strip() and p.placeholder_format.idx != 10:
                    p._element.getparent().remove(p._element)
        for i, sl in enumerate(self.prs.slides):
            if i == 0:
                continue
            for lp in sl.slide_layout.placeholders:
                if lp.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER:
                    sl.shapes._spTree.append(copy.deepcopy(lp._element))
        datei = Path(datei)
        datei.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(datei)
        print("gespeichert:", datei, len(self.prs.slides), "Folien")
        return datei
