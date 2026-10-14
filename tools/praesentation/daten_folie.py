"""Folie "Die Daten": welche Datensaetze, welcher Messort, wofuer genutzt.

Wird von projekt.py nach der Prototyp-Folie eingefuegt. Als Skript
aufgerufen, setzt es die Folie in eine bestehende (von Hand bearbeitete)
Projektpraesentation ein, ohne die uebrigen Folien zu veraendern:

    python tools/praesentation/daten_folie.py <ein.pptx> <aus.pptx> [position]
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from thm_deck import DUNKEL, GRAU, GRUEN, ROT, Deck  # noqa: E402

ZEILEN = [
    ["Datensatz", "Messort", "Stürze?", "Wofür im Projekt"],
    ["SisFall", "Hüfte", "ja (38 Pers.)", "Basismodell – Ausgangspunkt für alles"],
    ["KFall", "unterer Rücken", "ja (32)", "Zwischenstufe im Training"],
    ["UMAFall", "5 Stellen, u. a. Taille, Knöchel", "ja (19)", "erster Test am Knöchel"],
    ["FallAllD", "Nacken, Hand, Taille", "ja (15)", "Übertragbarkeit – am schwächsten"],
    ["WEDA-FALL", "Handgelenk", "ja (25, 77–95 J.)", "Fehlalarme bei Älteren"],
    ["CAPPIMU", "9 Stellen, u. a. Schienbein", "ja (30)", "Beinstürze fürs Fußmodell; links ≠ rechts"],
    ["ImbalanSense", "Sohle", "nein – nur Alltag (40)", "Training: senkt die Fehlalarme"],
    ["Smart-Insole", "Sohle", "nein – nur Alltag (29)", "nur Test, nie im Training"],
    ["Partnerdaten", "Sohle", "nein", "unbrauchbar: 10 Hz, ±2 g, ohne Gyroskop"],
    ["Eigene Pilotdaten", "Hüfte, Knöchel links/rechts", "ja (1 Pers.)", "Test auf eigener Hardware: 67 von 67"],
    ["Eigene Sohlenaufnahmen", "Sohle, beide Füße", "ja (1 Pers.)", "kein Sturz stärker als Joggen/Springen"],
]
SOHLE = (7, 8, 9, 11)


def daten_folie(d, s):
    """Fuellt die (leere Text-)Folie s."""
    t = d.tabelle(s, 0.47, 1.8, 12.3, ZEILEN, [2.45, 3.25, 2.3, 4.3], size=12, h_zeile=0.36,
                  zentriert_ab=4)
    for i in SOHLE:                      # Sohlen-Zeilen gruen markieren
        c = t.cell(i, 1)
        c.text_frame.paragraphs[0].runs[0].font.bold = True
        c.text_frame.paragraphs[0].runs[0].font.color.rgb = GRUEN
    d.text(s, 0.47, 6.2, 12.3, 0.75, [
        [("Kein öffentlicher Datensatz enthält Stürze an der Sohle. ", {"bold": True, "color": ROT}),
         ("Deshalb: Stürze vom Schienbein (CAPPIMU) + Alltag von der Sohle (ImbalanSense) → Modell B2r.",
          {"color": DUNKEL})]], size=15)
    d.notiz(s, "Stürze gibt es an Rumpf, Armen, Schienbein und Knöchel – an der Sohle nur Alltag. Die eigenen "
               "Sohlenaufnahmen enthalten Stürze, stammen aber von einer Person und dienen nur zum Prüfen. "
               "MobiAct (Smartphone, ±2 g, nicht kommerziell) wurde nicht verwendet.")
    return s


TITEL = "Die Daten: Stürze gibt es überall – nur nicht an der Sohle"

if __name__ == "__main__":
    from pptx import Presentation
    ein, aus = Path(sys.argv[1]), Path(sys.argv[2])
    pos = int(sys.argv[3]) if len(sys.argv) > 3 else 5          # 1-basiert
    d = Deck()
    d.prs = Presentation(ein)
    d.L = {l.name: l for l in d.prs.slide_layouts}
    s = d.folie(TITEL)
    daten_folie(d, s)
    # leere Platzhalter (ausser Foliennummer) entfernen
    for p in list(s.placeholders):
        if p.has_text_frame and not p.text_frame.text.strip() and p.placeholder_format.idx != 10:
            p._element.getparent().remove(p._element)
    import copy
    from pptx.enum.shapes import PP_PLACEHOLDER
    for lp in s.slide_layout.placeholders:
        if lp.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER:
            s.shapes._spTree.append(copy.deepcopy(lp._element))
    ids = d.prs.slides._sldIdLst
    neu = ids[-1]
    ids.remove(neu)
    ids.insert(pos - 1, neu)
    d.prs.save(aus)
    print("gespeichert:", aus, len(d.prs.slides), "Folien")
